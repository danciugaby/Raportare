#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Analiza evaluării activității didactice de către studenți – Universitatea Transilvania din Brașov.

Citește exportul platformei de evaluare (arhiva cu fișe PDF F05.1-PS8.2-01, rapoartele pe program
și tabelele *_SINTEZA.docx) plus, opțional, raportul de sinteză al facultății, apoi produce:

  rezultate/raport_evaluare.html   raportul vizual (se deschide în orice browser)
  rezultate/fise.csv               toate pozițiile cadru didactic × disciplină × C/A × program
  rezultate/comentarii.csv         toate comentariile studenților, cu teme etichetate automat
  rezultate/programe.csv           participare și scoruri pe program
  rezultate/dimensiuni.csv         niveluri S/M/R pe dimensiuni, curs vs aplicații
  rezultate/verificari.txt         reconcilieri și probleme găsite în date
  rezultate/date_complete.json     toate datele extrase, pentru alte analize

Utilizare:
  python analiza_evaluare.py --zip IESC.zip --raport-facultate Facultatea_....docx
  python analiza_evaluare.py --dir IESC --raport-facultate Facultatea_....docx --out rezultate

Fișierele CSV folosesc separatorul ';' și codificarea UTF-8 cu BOM, ca să se deschidă corect în Excel.
"""
from __future__ import annotations

import argparse
import collections
import csv
import html
import json
import logging
import re
import statistics
import sys
import unicodedata
import warnings
import zipfile
from multiprocessing import Pool
from pathlib import Path

warnings.filterwarnings("ignore")
logging.disable(logging.CRITICAL)

try:
    import pdfplumber
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
except ImportError as e:  # pragma: no cover
    sys.exit(f"Lipsește o bibliotecă ({e.name}). Rulați: pip install -r requirements.txt")

AICI = Path(__file__).resolve().parent

# ---------------------------------------------------------------------------
# Configurare
# ---------------------------------------------------------------------------

# Prescurtări pentru programe (cheia = numele programului fără diacritice, cu litere mici).
# Programele care nu apar aici primesc o prescurtare generată automat din inițiale.
PROGRAM_ALIAS = {
    "automatica si informatica aplicata": "AIA",
    "automatica si informatica aplicata 2026": "AIA 2026",
    "calculatoare": "C",
    "electronica aplicata": "EA",
    "electrotehnica": "ET",
    "inginerie electrica si calculatoare (in limba engleza)": "IEC-EN",
    "inginerie electronica, telecomunicatii si tehnologii informationale": "IETTI",
    "robotica": "ROB",
    "tehnologia informatiei": "TI",
    "tehnologii si sisteme de telecomunicatii": "TST",
    "securitate cibernetica(in limba engleza)": "SC-EN",
    "securitate cibernetica (in limba engleza)": "SC-EN",
    "sisteme avansate in automatica si tehnologii informatice": "SAATI",
    "sisteme electrice avansate (in limba engleza)": "SEA-EN",
    "sisteme electronice si de comunicatii integrate-rcd": "SECI-RCD",
    "sisteme electronice si de comunicatii integrate-si": "SECI-SI",
}

# Diacritice pentru numele departamentelor (în antetul rapoartelor apar fără diacritice).
DIACRITICE_DEP = {
    "Automatica": "Automatică", "Informatiei": "Informației", "Electronica": "Electronică",
    "Electrica": "Electrică", "Fizica": "Fizică", "Aplicata": "Aplicată", "Si": "și",
    "Inginerie": "Inginerie", "Tehnologia": "Tehnologia", "Calculatoare": "Calculatoare",
    "Mecanica": "Mecanică", "Telecomunicatii": "Telecomunicații",
}

DIM_NUME = [
    "Organizarea cursului/seminarului", "Strategii de instruire", "Mijloace și materiale didactice",
    "Relația profesor–student", "Calitatea învățării", "Aprecierea globală",
    "Implicarea studentului", "Respectarea criteriilor de evaluare",
]

# Teme pentru etichetarea automată a comentariilor (expresii regulate pe text fără diacritice).
TEME_NEG = {
    "Laborator/seminar fără explicații, studenți lăsați singuri": r"indrumar|lasa(ti)? singur|de capul nostru|fara ajutor|nu se explica|zero predare|nu ni s-a explicat|asteapta sa punem intrebari|autodidact|descurcam singuri|explicatii (mai|insuficiente|sa fie)|mai multe explicatii|sa explice|explicatiile .{0,20}insuficiente|copy.?paste",
    "Evaluare netransparentă sau inechitabilă": r"barem|notare|notarea|colocviu|depunct|obiectivitat|criterii\w* de notare|contest|feedback|evaluarea? (finala|de final|la)|examen\w* (a fost|nu are|fara)",
    "Conduită: limbaj jignitor, lipsă de respect": r"injur|vulgar|jign|umil|licentios|sinucid|omoram|asshole|misogyn|tente sexuale|irascib|nerespectuos|comportament neadecvat|arogant|ameninta|hate on|mean to|favourites|ridicand tonul|fucking|senile|bataie de joc|laserul",
    "Orar: ora 8, blocuri lungi, ore rare, sâmbătă": r"8 dimineata|ora 8|de la 8|la 8|dimineata|saturday|sambata|4 ore|mai multe ore|numarul de ore|un laborator pe saptamana|in fiecare saptamana|o data la (2|doua) saptamani|odata la doua|ora de incepere|ora pentru inceperea",
    "Prezența cadrului didactic, înlocuiri, atribuire": r"intarzi|online|nu a venit|sa fie prezent|mai prezent|laborant|nu apare|nu era in lista|evaluare (facuta )?pentru|punctualit",
    "Digresiuni fără legătură cu materia": r"politic|filozofie|divag|paranteze|tangent|se abate|vorbarie|nu au legatura|fara legatura|reproducer|concursuri|premiile|masini(lor)? de spalat|cotidian|abereaza|pe langa (tematica|subiect)",
    "Echipamente învechite sau defecte": r"aparatur|echipament|defect|invechit|placile|placutele|plc|reinnoite|modernizare",
    "Materiale cu erori sau incomplete": r"greseli|gresel|neactualizat|nu au fost actualizate|lipsuri|inexistent|non-existent|post (the )?answers|raspunsurile|suportul de laborator",
    "Volum disproporționat față de credite": r"credit|volum|dispropor|stresant|prea multa materie|1200",
}
TEME_POZ = {
    "Explicații clare, pe înțelesul tuturor": r"explic|inteles|intelesul|clar",
    "Activități practice, proiecte, exemple reale": r"practic|proiect|fpga|exemple|viata reala|real life|firme|companii|circuit|experiment|aplicatii|exercit|probleme|plc|automatele fizice|hands on|montaj",
    "Disponibilitate, răbdare, răspunsuri la întrebări": r"raspun|rabdar|ajut|sprijin|disponib|receptiv|intrebari",
    "Organizare și materiale bune": r"structur|organiz|suport de curs|materiale|resurse|bine redactat",
    "Atmosferă plăcută, relație bună cu studenții": r"atmosfer|prietenos|calm|glum|umor|relax|destins|amabil|chill|libera exprimare|fara sa ma critice",
    "Implicare și pasiune": r"implicat|pasionat|dedicat|implicare",
}
TRIVIAL = re.compile(
    r"^\W*(nimic\W*|nu|tot\W*|totul\W*|nothing|-+|\.+|x|:\)|:\(|numic|e ok.*|totul (este )?ok|"
    r"nu am nimic de (adaugat|aduagat|reprosat)\W*|nu am mentiuni\W*|n/a|momentan.*nimic.*|nimic de adaugat\W*|"
    r"nu stiu|habar n am|nimic diferit|all of it|super|e bun asa cum este|a fost in regula\W*|nu am obiectiuni|"
    r"nimic[,.]? (este|e) (in regula|bine|perfect).*|cam nimic)?\W*$")
NEGATIV_IN_PLACUT = re.compile(r"^\W*(nimic|nu mi-?a? ?placut|nu mi a placut|nu am nimic)|worst|bataie de joc|irascib|iritabil|vulgar|dezinteresat")
LEAK_EXAMEN = re.compile(r"\s*VIII\.8\..*$", re.S)

# ---------------------------------------------------------------------------
# Utilitare
# ---------------------------------------------------------------------------

def norm(s: str | None) -> str:
    """Cheie de comparare: fără diacritice, majuscule, doar litere și cifre."""
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c)).upper()
    return re.sub(r"[^A-Z0-9]", "", s)


def fara_diacritice(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    return "".join(c for c in s if not unicodedata.combining(c)).lower().replace("?", "")


def repara_mojibake(s: str) -> str:
    """Unele nume din arhivă au literele Ă/Ț/Î stocate ca simboluri cp437 (╞, ▌, ╫). Le refacem prin cp852."""
    out = []
    for c in s:
        if 0x2500 <= ord(c) <= 0x25FF:
            try:
                out.append(bytes([c.encode("cp437")[0]]).decode("cp852"))
                continue
            except (UnicodeEncodeError, UnicodeDecodeError):
                pass
        out.append(c)
    return "".join(out)


def cod_program(nume_fisier: str) -> str | None:
    m = re.search(r"\((\d+)\)(?:_SINTEZA)?\.(pdf|docx)$", nume_fisier, re.I)
    return m.group(1) if m else None


def pct(a, b) -> int:
    return round(100 * a / b) if b else 0


def f2(x, d=2) -> str:
    return "—" if x is None else f"{x:.{d}f}".replace(".", ",")


def medie_ponderata(recs) -> float | None:
    rs = [r for r in recs if r["n17"] and r["score17"] is not None]
    n = sum(r["n17"] for r in rs)
    return round(sum(r["score17"] * r["n17"] for r in rs) / n, 2) if n else None


# ---------------------------------------------------------------------------
# 1. Arhiva
# ---------------------------------------------------------------------------

def extrage_arhiva(zip_path: Path, dest: Path) -> tuple[Path, int]:
    """Extrage arhiva refăcând numele de fișiere (UTF-8 fără flag, cp852, simboluri cp437)."""
    dest.mkdir(parents=True, exist_ok=True)
    corupte = 0
    with zipfile.ZipFile(zip_path) as z:
        for info in z.infolist():
            nume = info.filename
            if not info.flag_bits & 0x800:
                brut = nume.encode("cp437")
                try:
                    nume = brut.decode("utf-8")
                except UnicodeDecodeError:
                    nume = brut.decode("cp852")
            reparat = repara_mojibake(nume)
            if reparat != nume:
                corupte += 1
            if reparat.endswith("/"):
                continue
            tinta = dest / reparat
            tinta.parent.mkdir(parents=True, exist_ok=True)
            tinta.write_bytes(z.read(info))
    return dest, corupte


def gaseste_fisiere(root: Path):
    pdfs = sorted(p for p in root.rglob("*") if p.suffix.lower() == ".pdf")
    sinteze = sorted(p for p in root.rglob("*_SINTEZA.docx"))
    rapoarte = sorted(p for p in root.rglob("*.docx") if not p.name.endswith("_SINTEZA.docx") and cod_program(p.name))
    return pdfs, rapoarte, sinteze


# ---------------------------------------------------------------------------
# 2. Fișele PDF
# ---------------------------------------------------------------------------

NIVELURI = ("SCĂZUT", "MEDIU", "RIDICAT")
ETICHETE_ANTET = {"Cadru didactic:": "cd", "Disciplina:": "disc", "Program de studii:": "prog",
                  "An de studiu:": "an", "Data evaluării:": "data"}


def citeste_fisa(path_str: str) -> dict:
    path = Path(path_str)
    nume = repara_mojibake(path.name)
    r = {"file": nume, "prog_code": cod_program(nume),
         "tip": "C" if " Curs(" in nume else ("A" if " Aplicatii(" in nume else "?"),
         "levels": [], "likes": [], "dislikes": [], "blank_pages": 0, "html_tags": False}
    try:
        pdf = pdfplumber.open(path)
    except Exception as e:  # fișier deteriorat
        r["error"] = str(e)
        r.update(n17=0, n8=0, score17=None, score8=None, pages=0)
        return r
    with pdf:
        r["pages"] = len(pdf.pages)
        # antet: etichete în stânga, valori la x≈208
        words = pdf.pages[0].extract_words(keep_blank_chars=True, x_tolerance=2)
        labs = [(w["top"], ETICHETE_ANTET[w["text"].strip()]) for w in words
                if w["text"].strip() in ETICHETE_ANTET and w["x0"] < 100]
        for w in words:
            if not (190 < w["x0"] < 260 and w["top"] < 140):
                continue
            t = w["text"].strip()
            best = min(labs, key=lambda l: abs(l[0] - w["top"])) if labs else None
            if best and abs(best[0] - w["top"]) < 4:
                r[best[1]] = t
            elif t in ("Curs", "Aplicatii", "Aplicații", "Seminar", "Laborator", "Proiect"):
                r["tip_txt"] = t
            elif re.match(r"\d+/\d+/\d{4}", t):
                r["data"] = t
        lines = []
        for i, pg in enumerate(pdf.pages):
            tl = pg.extract_text_lines()
            if len("".join(l["text"] for l in tl if not l["text"].startswith("F05.1"))) < 8:
                r["blank_pages"] += 1
            lines += [(i, l["top"], l["x0"], l["text"]) for l in tl]
            # nivelul bifat (simbol Wingdings) pe fiecare dimensiune
            checks = [c for c in pg.chars if "Wingdings" in c.get("fontname", "") and c["text"].strip()]
            if checks:
                ws = pg.extract_words()
                for c in checks:
                    rand = [w for w in ws if abs(w["top"] - c["top"]) < 4 and w["text"] in NIVELURI]
                    stanga = [w for w in rand if w["x0"] < c["x0"]]
                    nivel = max(stanga, key=lambda w: w["x0"])["text"] if stanga else None
                    titluri = [l for l in tl if re.match(r"^\s*\d\s*\.", l["text"]) and c["top"] - 40 <= l["top"] <= c["top"] + 2]
                    dim = int(re.match(r"^\s*(\d)", max(titluri, key=lambda l: l["top"])["text"]).group(1)) if titluri else None
                    r["levels"].append((dim, nivel))
    text = "\n".join(t for *_, t in lines)
    m = re.search(r"Numar total de respondenti D1-D7:\s*(\d+)", text); r["n17"] = int(m.group(1)) if m else 0
    m = re.search(r"Numar total de respondenti D8:\s*(\d+)", text); r["n8"] = int(m.group(1)) if m else 0
    m = re.search(r"Scor total D1-D7\s*([\d.,]+)?\s*Pentru", text); r["score17"] = float(m.group(1).replace(",", ".")) if m and m.group(1) else None
    m = re.search(r"Scor total D8\s*([\d.,]+)?\s*Pentru", text); r["score8"] = float(m.group(1).replace(",", ".")) if m and m.group(1) else None
    # comentarii
    stare = None
    for _, _, x0, t in lines:
        if t.startswith("Ce v-a plăcut"):
            stare = "likes"; r["html_tags"] |= "<br/>" in t; continue
        if t.startswith(("Ce ați vrea", "Ce aţi vrea")):
            stare = "dislikes"; continue
        if t.startswith("Concluzii"):
            stare = None; continue
        if stare is None or t.startswith("F05.1-PS8.2") or re.fullmatch(r"\s*\d+\s*", t):
            continue
        if "features.)" in t or "this course?" in t or t.startswith("3 features"):
            continue
        m = re.match(r"^(\d+)\.\s?(.*)$", t)
        if m and x0 < 80:
            r[stare].append(m.group(2).strip())
        elif r[stare]:
            r[stare][-1] += " " + t.strip()
    return r


# ---------------------------------------------------------------------------
# 3. Rapoartele și tabelele de sinteză (docx)
# ---------------------------------------------------------------------------

def _randuri(t, acc):
    for row in t.rows:
        cells, seen = [], set()
        for c in row.cells:
            if id(c._tc) in seen:
                continue
            seen.add(id(c._tc))
            if c.tables:
                acc.append(("P", c.text.strip()))
                for nt in c.tables:
                    _randuri(nt, acc)
            else:
                cells.append(c.text.strip())
        if any(cells):
            acc.append(("R", cells))


def _incarca(path: Path):
    d = docx.Document(str(path))
    acc, are_grafice_sIII, in_sIII = [], False, False
    for el in d.element.body.iterchildren():
        txt = "".join(x.text or "" for x in el.iter(qn("w:t")))
        if "III. PREZENTAREA GRAFICĂ" in txt:
            in_sIII = True
        if "IV. CONCLUZII" in txt:
            in_sIII = False
        if in_sIII and (el.findall(".//" + qn("w:drawing")) or el.findall(".//" + qn("w:pict"))):
            are_grafice_sIII = True
        if el.tag.endswith("}tbl"):
            _randuri(Table(el, d), acc)
        elif el.tag.endswith("}p") and txt.strip():
            acc.append(("P", txt.strip()))
    return acc, are_grafice_sIII


def citeste_raport(path: Path) -> dict:
    acc, grafice = _incarca(path)
    out = {"file": repara_mojibake(path.name), "code": cod_program(path.name), "grafice_sIII": grafice}
    tot = " ".join(x[1] if x[0] == "P" else " | ".join(x[1]) for x in acc)
    out["text"] = tot
    m = re.search(r"Programul de studii:\s*(.*?)\s+I\. INTRODUCERE", tot); out["program"] = (m.group(1).strip() or None) if m else None
    m = re.search(r"An universitar:\s*(.*?)\s*/?\s*/?\s*Programul de studii", tot); out["an_univ"] = (m.group(1).strip(" /") or None) if m else None
    m = re.search(r"DEPARTAMENTUL\s+([A-ZĂÂÎȘȚŞŢ ,\-]+?)\s+RAPORT", tot); out["departament"] = m.group(1).strip() if m else None
    out["are_facultate_in_antet"] = "FACULTATEA" in tot.split("RAPORT")[0]
    m = re.search(r"aplicat unui număr de\s*(\d+)\s*studenți", tot); out["studenti"] = int(m.group(1)) if m else None
    m = re.search(r"obținut\s*(\d+)\s*răspunsuri la evaluarea pe parcursul semestrului și\s*(\d+)", tot)
    out["resp_sem"], out["resp_exam"] = (int(m.group(1)), int(m.group(2))) if m else (None, None)
    R = [x[1] for x in acc if x[0] == "R"]
    numerice = [r for r in R if len(r) == 8 and all(re.match(r"^[\d.]+$", c) for c in r)]
    out["t11"] = [float(c) for c in numerice[0]] if len(numerice) >= 1 else None
    out["t12"] = [float(c) for c in numerice[1]] if len(numerice) >= 2 else None
    rows, grp = [], None
    for r in R:
        if r and r[0] in ("C", "A") and len(r) >= 2:
            if len(r) >= 9 and re.match(r"^\d+$", r[1]):
                rows.append({"tip": r[0], "an": r[3], "disc": r[4], "cd": r[5], "nstud": int(r[6] or 0),
                             "sem": int(r[7] or 0), "exam": int(r[8] or 0), "grp": grp})
            elif r[1].startswith(("Licenta", "Masterat", "Doctorat")):
                grp = r[1]
            elif "Total" in " ".join(r):
                out["total_" + r[0]] = r
        elif r and "TOTAL GENERAL" in " ".join(r):
            out["total_general"] = r
    out["rows"] = rows

    def num(x):
        try:
            return float(x)
        except ValueError:
            return None
    dims = [r for r in R if len(r) == 4 and re.match(r"^\d\.", r[0])]
    out["t2"] = [[num(c) for c in r[1:]] for r in dims[:8]] if len(dims) >= 16 else None
    out["t3"] = [[num(c) for c in r[1:]] for r in dims[8:16]] if len(dims) >= 16 else None
    out["zero_in_t23"] = sum(1 for r in dims for c in r[1:] if c.strip() in ("0.00", "0,00"))
    out["perioada_goala"] = "perioada ……" in tot or "perioada ..." in tot
    out["data_aprobare_goala"] = bool(re.search(r"din data de\s*[.…]{3,}", tot))
    m = re.search(r"măsurile de îmbunătățire propuse\.(.*?)Raportul a fost aprobat", tot, re.S)
    out["concluzii_goale"] = bool(m) and not re.sub(r"[\s/]", "", m.group(1))
    out["tabel2_de_doua_ori"] = tot.count("este prezentată în tabelul 2") >= 2
    out["tipografie"] = [s for s in ("2s-a", "tabelele1.1") if s in tot]
    out["nan"] = "NaN" in tot
    return out


def _grila(tbl):
    out = []
    for tr in tbl.findall(qn("w:tr")):
        row = []
        for tc in tr.findall(qn("w:tc")):
            pr = tc.find(qn("w:tcPr")); gs = 1
            if pr is not None and pr.find(qn("w:gridSpan")) is not None:
                gs = int(pr.find(qn("w:gridSpan")).get(qn("w:val")))
            row += ["".join(x.text or "" for x in tc.iter(qn("w:t"))).strip()] + [""] * (gs - 1)
        out.append(row)
    return out


def citeste_sinteza(path: Path) -> dict:
    d = docx.Document(str(path))
    texte = [t for t in ("".join(x.text or "" for x in p.iter(qn("w:t"))).strip() for p in d.element.body.iter(qn("w:p"))) if t]
    out = {"file": repara_mojibake(path.name), "code": cod_program(path.name), "header": texte[:8],
           "nivel": next((t for t in texte if t.upper().startswith(("LICENTA", "MASTERAT", "DOCTORAT"))), ""), "rows": []}
    for tbl in d.element.body.iter(qn("w:tbl")):
        g = _grila(tbl)
        hi = [i for i, r in enumerate(g) if r and r[0] == "An de studii"]
        if not hi:
            continue
        h1, h2 = g[hi[0]], g[hi[0] + 1]
        etich, cur = [], None
        for a, b in zip(h1, h2):  # coloanele S/M/R lipsesc când sunt goale peste tot – le citim din antet
            if re.match(r"^D\d$", a):
                cur = a
            etich.append((cur, b) if b in ("S", "M", "R") else (None, None))
        an = disc = tip = None
        for r in g[hi[0] + 2:]:
            if len(r) < 5 or not re.match(r"^\d+\s*/\s*\d+$", r[4] or ""):
                continue
            if r[0]: an = r[0]
            if r[1]: disc, tip = r[1], None
            if r[2]: tip = r[2]
            n17, n8 = (int(x) for x in r[4].split("/"))
            dims = {}
            for (dl, lv), v in zip(etich, r):
                if dl:
                    dims.setdefault(dl, {"S": 0, "M": 0, "R": 0})[lv] = int(v) if v.isdigit() else 0
            out["rows"].append({"an": an, "disc": disc, "tip": tip, "cd": r[3], "n17": n17, "n8": n8, "dims": dims})
    return out


# ---------------------------------------------------------------------------
# 4. Îmbinarea surselor
# ---------------------------------------------------------------------------

def separa_nume(baza: str, cunoscute: list[str]) -> tuple[str, str]:
    """Separă 'NUME Prenume Disciplina' din numele fișierului unei fișe goale."""
    nb = norm(baza)
    cuvinte = [w for w in baza.split(" ") if w]
    cdn = next((n for n in cunoscute if nb.startswith(n)), None)
    if cdn:
        for k in range(1, len(cuvinte)):
            if norm(" ".join(cuvinte[:k])) == cdn:
                return " ".join(cuvinte[:k]), " ".join(cuvinte[k:])
    j = next((i for i, w in enumerate(cuvinte) if i > 0 and w.islower()), None)
    if j is not None and j >= 2:
        k = j - 1
    elif cuvinte and cuvinte[0].isupper():
        k = 1
        while k < len(cuvinte) - 1 and cuvinte[k].isupper():
            k += 1
    else:
        k = 2 if len(cuvinte) >= 3 else 1
    return " ".join(cuvinte[:k]), " ".join(cuvinte[k:])


def imbina(fise: list[dict], rapoarte: list[dict], sinteze: list[dict]):
    sin = [dict(x, prog=s["code"]) for s in sinteze for x in s["rows"]]
    cunoscute = sorted({norm(x["cd"]) for x in sin} | {norm(r.get("cd")) for r in fise if r.get("cd")} - {""}, key=len, reverse=True)
    recs = []
    for r in fise:
        baza = re.sub(r" (Curs|Aplicatii)\(\d+\)\.pdf$", "", r["file"])
        if r.get("cd") and r.get("disc"):
            cd, disc = r["cd"], r["disc"]
        else:
            cd, disc = separa_nume(baza, cunoscute)
        recs.append(dict(r, cd_disp=cd, disc_disp=disc, cdn=norm(cd), discn=norm(disc)))
    idx = collections.defaultdict(list)
    for i, r in enumerate(recs):
        idx[(r["prog_code"], r["tip"], r["cdn"], r["discn"])].append(i)
    lipsa = []
    for s in sin:
        k = (s["prog"], s["tip"], norm(s["cd"]), norm(s["disc"]))
        cand = [i for i in idx.get(k, []) if "sin" not in recs[i] and recs[i]["n17"] == s["n17"] and recs[i]["n8"] == s["n8"]]
        if cand:
            recs[cand[0]]["sin"] = s
            recs[cand[0]]["an"] = recs[cand[0]].get("an") or s["an"]
        else:  # rând din sinteză fără fișă PDF (de ex. fișă suprascrisă în arhivă)
            lipsa.append(s)
            recs.append(dict(file=None, prog_code=s["prog"], tip=s["tip"], cd_disp=s["cd"], disc_disp=s["disc"],
                             cdn=norm(s["cd"]), discn=norm(s["disc"]), an=s["an"], n17=s["n17"], n8=s["n8"],
                             score17=None, score8=None, likes=[], dislikes=[], levels=[], missing_pdf=True, sin=s))
    fara_sinteza = [r["file"] for r in recs if "sin" not in r and (r["n17"] or r["n8"])]
    return recs, lipsa, fara_sinteza


def prescurtare(program: str) -> str:
    cheie = fara_diacritice(program).strip()
    if cheie in PROGRAM_ALIAS:
        return PROGRAM_ALIAS[cheie]
    en = "-EN" if "engleza" in cheie else ""
    cheie = re.sub(r"\(.*?\)", "", cheie)
    an = re.search(r"\b(20\d\d)\b", cheie)
    sufix = re.search(r"-([a-z]+)$", cheie.strip())
    cuv = [w for w in re.split(r"[\s,\-]+", re.sub(r"-[a-z]+$", "", cheie)) if w and w not in ("si", "de", "in", "a", "al", "pentru") and not w.isdigit()]
    s = "".join(w[0].upper() for w in cuv) or program[:3].upper()
    if len(s) == 1:
        s = cuv[0][:3].upper()
    return s + (("-" + sufix.group(1).upper()) if sufix else "") + en + ((" " + an.group(1)) if an else "")


def nume_departament(s: str | None) -> str:
    if not s:
        return "—"
    return " ".join(DIACRITICE_DEP.get(w.title(), w.title()) for w in s.split())


def initiale_dep(s: str | None) -> str:
    if not s:
        return "—"
    return "".join(w[0] for w in s.split() if w.upper() not in ("SI", "DE", "A"))


# ---------------------------------------------------------------------------
# 5. Analize
# ---------------------------------------------------------------------------

def niveluri(recs, D):
    c = collections.Counter()
    for r in recs:
        if "sin" in r and D in r["sin"]["dims"]:
            c.update(r["sin"]["dims"][D])
    return {"S": c["S"], "M": c["M"], "R": c["R"], "n": sum(c.values())}


def teme(recs):
    """Etichetare automată: în câte fișe apare fiecare temă (o fișă poate intra la mai multe)."""
    neg = {k: set() for k in TEME_NEG}
    poz = {k: set() for k in TEME_POZ}
    randuri = []
    subst_neg = subst_poz = 0
    for i, r in enumerate(recs):
        for rubrica, lista in (("plăcut", r["likes"]), ("altfel", r["dislikes"])):
            for t in lista:
                curat = LEAK_EXAMEN.sub("", t).strip()
                s = fara_diacritice(curat)
                trivial = not s or bool(TRIVIAL.match(s))
                etich = []
                if not trivial:
                    if rubrica == "plăcut" and not NEGATIV_IN_PLACUT.search(s):
                        subst_poz += 1
                        for k, p in TEME_POZ.items():
                            if re.search(p, s):
                                poz[k].add(i); etich.append(k)
                    else:
                        subst_neg += 1
                        for k, p in TEME_NEG.items():
                            if re.search(p, s):
                                neg[k].add(i); etich.append(k)
                randuri.append((i, rubrica, curat, trivial, etich, "VIII.8" in t))
    return ({k: len(v) for k, v in neg.items()}, {k: len(v) for k, v in poz.items()}, randuri, subst_neg, subst_poz)


def teme_manuale(csv_path: Path, recs):
    """Numără temele din coloana 'tema_manuala' a unui comentarii.csv completat de mână (teme separate prin '|')."""
    neg, poz = collections.defaultdict(set), collections.defaultdict(set)
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f, delimiter=";"):
            for t in (row.get("tema_manuala") or "").split("|"):
                t = t.strip()
                if t:
                    (poz if row.get("rubrica") == "plăcut" else neg)[t].add(row.get("id_fisa"))
    return {k: len(v) for k, v in neg.items()}, {k: len(v) for k, v in poz.items()}


def analizeaza(recs, rapoarte, sinteze, fac, prag_n, prag_scor, prag_s):
    rep = {r["code"]: r for r in rapoarte}
    sin_meta = {s["code"]: s for s in sinteze}
    coduri = sorted({r["prog_code"] for r in recs} | set(rep), key=lambda c: (rep.get(c, {}).get("departament") or "", c))
    progs = []
    for c in coduri:
        p = rep.get(c, {})
        rs = [r for r in recs if r["prog_code"] == c]
        nume = p.get("program") or (p.get("file", c).split("...")[0].title() if p else c)
        nivel = (sin_meta.get(c, {}).get("nivel") or "").upper()
        if not nivel and p.get("rows"):
            nivel = (p["rows"][0]["grp"] or "").upper()
        d6 = niveluri(rs, "D6")
        t11 = p.get("t11")
        progs.append(dict(code=c, short=prescurtare(nume), name=nume, niv="M" if "MASTER" in nivel else ("L" if "LICENTA" in nivel else ""),
                          dep=initiale_dep(p.get("departament")), depName=nume_departament(p.get("departament")),
                          stud=p.get("studenti") or 0, sem=p.get("resp_sem") or 0, exam=p.get("resp_exam") or 0,
                          sheets=len(rs), withAny=sum(1 for r in rs if r["n17"] or r["n8"]), withSem=sum(1 for r in rs if r["n17"]),
                          repDEv=[t11[4], t11[5]] if t11 else None,
                          score=medie_ponderata(rs), scoreC=medie_ponderata([r for r in rs if r["tip"] == "C"]),
                          scoreA=medie_ponderata([r for r in rs if r["tip"] == "A"]), d6=d6))
    progs_date = [p for p in progs if p["sheets"]]
    short = {p["code"]: p["short"] for p in progs}
    depcod = {p["code"]: p["dep"] for p in progs}
    deps = []
    for dcod in sorted({p["dep"] for p in progs_date}):
        cs = [p["code"] for p in progs_date if p["dep"] == dcod]
        rs = [r for r in recs if r["prog_code"] in cs]
        deps.append(dict(id=dcod, name=next(p["depName"] for p in progs_date if p["dep"] == dcod),
                         progs=[short[c] for c in cs], stud=sum(p["stud"] for p in progs_date if p["dep"] == dcod),
                         sem=sum(p["sem"] for p in progs_date if p["dep"] == dcod), exam=sum(p["exam"] for p in progs_date if p["dep"] == dcod),
                         sheets=len(rs), withAny=sum(1 for r in rs if r["n17"] or r["n8"]), score=medie_ponderata(rs),
                         scoreC=medie_ponderata([r for r in rs if r["tip"] == "C"]), scoreA=medie_ponderata([r for r in rs if r["tip"] == "A"]),
                         d6=niveluri(rs, "D6")))
    # media pe dimensiune: număr respondenți pe nivel × media nivelului (tabelele 2–3)
    dims = []
    for k in range(8):
        D = f"D{k + 1}"
        row = {"id": D, "name": DIM_NUME[k]}
        for t, tabel in (("C", "t2"), ("A", "t3")):
            l = niveluri([r for r in recs if r["tip"] == t], D)
            if fac and fac.get(tabel):
                m = dict(zip("SMR", fac[tabel][k]))
                est = sum(l[x] * m[x] for x in "SMR" if m[x]) / l["n"] if l["n"] else None
            else:  # fără raportul facultății: mediile pe nivel din rapoartele fiecărui program
                s = n = 0
                for p in rapoarte:
                    if not p.get(tabel):
                        continue
                    lp = niveluri([r for r in recs if r["tip"] == t and r["prog_code"] == p["code"]], D)
                    for i, x in enumerate("SMR"):
                        if p[tabel][k][i]:
                            s += lp[x] * p[tabel][k][i]; n += lp[x]
                est = s / n if n else None
            row[t] = dict(l, est=round(est, 2) if est is not None else None)
        dims.append(row)
    # pe fișă
    sheets = []
    for i, r in enumerate(recs):
        d6 = (r.get("sin") or {}).get("dims", {}).get("D6", {})
        nd = sum(d6.values())
        sheets.append(dict(id=i, p=short.get(r["prog_code"], r["prog_code"]), dep=depcod.get(r["prog_code"], "—"), t=r["tip"],
                           cd=r["cd_disp"], d=r["disc_disp"], an=r.get("an") or "", n=r["n17"], n8=r["n8"], s=r["score17"],
                           s8=r["score8"], sS=round(100 * d6.get("S", 0) / nd) if nd else None, miss=bool(r.get("missing_pdf")),
                           nc=len(r["likes"]) + len(r["dislikes"])))
    ns = sorted(r["n17"] for r in recs if r["n17"])
    depth = []
    for lo, hi, l in [(1, 1, "1"), (2, 2, "2"), (3, 4, "3–4"), (5, 9, "5–9"), (10, 19, "10–19"), (20, 10**6, "20+")]:
        k = [x for x in ns if lo <= x <= hi]
        depth.append(dict(l=l, sheets=len(k), resp=sum(k)))
    cu_scor = [r for r in recs if r["n17"] and r["score17"] is not None]
    hist = []
    for lo, hi in [(1, 1.5), (1.5, 2), (2, 2.5), (2.5, 3), (3, 3.5), (3.5, 4), (4, 4.5), (4.5, 5.01)]:
        hist.append(dict(lo=lo, hi=min(hi, 5), k=sum(1 for r in cu_scor if lo <= r["score17"] < hi),
                         k5=sum(1 for r in cu_scor if r["n17"] >= 5 and lo <= r["score17"] < hi)))
    # acoperire reală pe discipline
    def acop(t):
        toate = {(r["prog_code"], r["discn"]) for r in recs if r["tip"] == t}
        sem = {(r["prog_code"], r["discn"]) for r in recs if r["tip"] == t and r["n17"]}
        return pct(len(sem), len(toate)), len(sem), len(toate)
    nstud = {t: sum(x["nstud"] for p in rapoarte for x in p["rows"] if x["tip"] == t) for t in "CA"}
    sem_t = {t: sum(r["n17"] for r in recs if r["tip"] == t) for t in "CA"}
    tot_lv = collections.Counter()
    for r in recs:
        if "sin" in r:
            for k in range(1, 8):
                tot_lv.update(r["sin"]["dims"].get(f"D{k}", {}))
    tot = dict(
        stud=(fac or {}).get("studenti") or sum(p["stud"] for p in progs),
        sem=(fac or {}).get("resp_sem") or sum(p["sem"] for p in progs),
        exam=(fac or {}).get("resp_exam") or sum(p["exam"] for p in progs),
        sheets=len(recs), withAny=sum(1 for r in recs if r["n17"] or r["n8"]), withSem=sum(1 for r in recs if r["n17"]),
        withExam=sum(1 for r in recs if r["n8"]), exam1=sum(1 for r in recs if r["n8"] == 1),
        empty=sum(1 for r in recs if not r["n17"] and not r["n8"]),
        teachers=len({r["cdn"] for r in recs}), teachAny=len({r["cdn"] for r in recs if r["n17"] or r["n8"]}),
        median=statistics.median(ns) if ns else 0, n12=sum(1 for x in ns if x <= 2),
        score=medie_ponderata(recs), scoreC=medie_ponderata([r for r in recs if r["tip"] == "C"]),
        scoreA=medie_ponderata([r for r in recs if r["tip"] == "A"]),
        withScore=len(cu_scor), over45=sum(1 for r in cu_scor if r["score17"] >= 4.5),
        lowSmall=sum(1 for r in cu_scor if r["n17"] < prag_n and r["score17"] < 3),
        shareR=pct(tot_lv["R"], sum(tot_lv.values())),
        rateC=round(100 * sem_t["C"] / nstud["C"], 1) if nstud["C"] else None,
        rateA=round(100 * sem_t["A"] / nstud["A"], 1) if nstud["A"] else None,
        nstud=nstud["C"] + nstud["A"], covC=acop("C"), covA=acop("A"),
        d8=round(sum(r["score8"] * r["n8"] for r in recs if r["n8"] and r["score8"] is not None) /
                 max(1, sum(r["n8"] for r in recs if r["n8"] and r["score8"] is not None)), 2),
        lowExamProgs=[p["short"] for p in progs_date if p["exam"] <= 4 and p["sem"] > 0],
        facDEv=[(fac or {}).get("t11", [None] * 8)[4], (fac or {}).get("t11", [None] * 8)[5]] if fac and fac.get("t11") else None,
        facCDEv=[fac["t11"][6], fac["t11"][7]] if fac and fac.get("t11") else None,
    )
    prag = dict(n=prag_n, scor=prag_scor, s=prag_s)
    return dict(progs=progs_date, deps=deps, dims=dims, sheets=sheets, depth=depth, hist=hist, tot=tot, prag=prag)


def verifica(recs, rapoarte, sinteze, fac, lipsa, fara_sinteza, nume_corupte, randuri_com):
    """Reconcilieri și probleme de date. Returnează (reconcilieri, probleme) ca liste de texte."""
    rec, prob = [], collections.defaultdict(list)
    sem_rap = sum(p["resp_sem"] or 0 for p in rapoarte); ex_rap = sum(p["resp_exam"] or 0 for p in rapoarte)
    sem_t13 = sum(x["sem"] for p in rapoarte for x in p["rows"]); ex_t13 = sum(x["exam"] for p in rapoarte for x in p["rows"])
    sem_sin = sum(x["n17"] for s in sinteze for x in s["rows"]); ex_sin = sum(x["n8"] for s in sinteze for x in s["rows"])
    sem_pdf = sum(r["n17"] for r in recs if r.get("file")); ex_pdf = sum(r["n8"] for r in recs if r.get("file"))
    rec.append(f"Răspunsuri semestru / sesiune — text rapoarte program: {sem_rap}/{ex_rap}; tabele 1.3: {sem_t13}/{ex_t13}; "
               f"tabele sinteză: {sem_sin}/{ex_sin}; fișe PDF: {sem_pdf}/{ex_pdf}" + (f"; raport facultate: {fac['resp_sem']}/{fac['resp_exam']}" if fac else ""))
    rec.append(f"Rânduri: tabele 1.3 pe program {sum(len(p['rows']) for p in rapoarte)}; tabele sinteză {sum(len(s['rows']) for s in sinteze)}"
               + (f"; tabel 1.3 facultate {len(fac['rows'])}" if fac else ""))
    if fac:
        st = sum(p["studenti"] or 0 for p in rapoarte)
        st_date = sum(p["studenti"] or 0 for p in rapoarte if p["rows"])
        rec.append(f"Studenți: raport facultate {fac['studenti']}; suma programelor {st} (din care programe fără nicio evaluare: {st - st_date})")
        F = prob["Raportul facultății"]
        goale = [n for n, v in (("an universitar", not fac.get("an_univ")), ("programul de studii", not fac.get("program")),
                                 ("perioada evaluării", fac["perioada_goala"]), ("data aprobării", fac["data_aprobare_goala"])) if v]
        if not (fac.get("total_C") or ["", "", "", "", "", "", ""])[6]:
            goale.append("totalul „Nr. Stud” în tabelul 1.3")
        if goale:
            F.append("Câmpuri necompletate: " + ", ".join(goale) + ".")
        if not fac["grafice_sIII"]:
            F.append("Secțiunea III „Prezentarea grafică” nu conține niciun grafic.")
        if fac["concluzii_goale"]:
            F.append("Secțiunea IV „Concluzii și măsuri de îmbunătățire” este goală.")
        if not fac["are_facultate_in_antet"]:
            F.append("Antetul nu conține numele facultății.")
        if fac["tabel2_de_doua_ori"]:
            F.append("Textul trimite la „tabelul 2” și pentru aplicații (ar trebui „tabelul 3”).")
        if fac["tipografie"]:
            F.append("Greșeli de tehnoredactare: " + ", ".join(f"„{t}”" for t in fac["tipografie"]) + ".")
        if fac["studenti"] and fac["studenti"] > st_date:
            F.append(f"Numărul de studenți ({fac['studenti']}) include {fac['studenti'] - st_date} studenți din programe fără nicio evaluare.")
    M = prob["Metoda de calcul a indicatorilor"]
    if fac and fac.get("t11"):
        M.append(f"„Studenți care au completat chestionare (%)” = răspunsuri / studenți ({f2(fac['t11'][0])}% = răspunsuri la curs / {fac['studenti']}); nu măsoară studenți distincți.")
        nC = sum(x["nstud"] for x in fac["rows"] if x["tip"] == "C"); nA = sum(x["nstud"] for x in fac["rows"] if x["tip"] == "A")
        sC = sum(x["sem"] for x in fac["rows"] if x["tip"] == "C"); sA = sum(x["sem"] for x in fac["rows"] if x["tip"] == "A")
        if nC and nA:
            M.append(f"„Chestionare completate (%)” împarte răspunsurile la curs la totalul curs + aplicații ({sC}/{nC + nA}). "
                     f"Calculat separat: curs {f2(100 * sC / nC, 1)}%, aplicații {f2(100 * sA / nA, 1)}%.")
    M.append("DPEv și CDPEv numără doar disciplinele și cadrele didactice cu cel puțin un răspuns, nu întreg statul de funcții, deci DEv% și CDEv% ies umflate.")
    z = sum(p["zero_in_t23"] for p in rapoarte)
    if z:
        M.append(f"În tabelele 2–3 pe program apare 0,00 de {z} ori, acolo unde nu există respondenți în acel nivel.")
    M.append("Tabelele 2–3 nu arată câți respondenți sunt în fiecare nivel.")
    nan = [p["file"] for p in rapoarte if p["nan"]]
    if nan:
        M.append(f"{len(nan)} rapoarte de program sunt goale (valori NaN): " + ", ".join(nan) + ".")
    P = prob["Fișele individuale (PDF)"]
    goale = [r for r in recs if r.get("file") and not r["n17"] and not r["n8"]]
    eticheta = sum(1 for r in goale if r["tip"] == "C" and r.get("tip_txt") == "Aplicatii")
    P.append(f"{len(goale)} fișe goale (fără niciun răspuns); antetul nu conține numele cadrului didactic și al disciplinei"
             + (f", iar {eticheta} fișe de curs goale poartă eticheta „Aplicatii”." if eticheta else "."))
    for s in lipsa:
        P.append(f"Lipsește fișa PDF: {s['cd']}, „{s['disc']}”, anul {s['an']}, {'curs' if s['tip'] == 'C' else 'aplicații'} "
                 f"(cod {s['prog']}, {s['n17']} răspunsuri). Probabil suprascrisă: numele fișierului nu conține anul de studiu.")
    for f in fara_sinteza:
        P.append(f"Fișă cu răspunsuri care nu apare în tabelul de sinteză: {f}.")
    intreb = sum(1 for i, rub, t, *_ in randuri_com if re.search(r"\w\?\w", t))
    if intreb:
        P.append(f"În {intreb} comentarii literele ș, ț, ă apar ca „?” (problemă de codificare).")
    tags = sum(1 for r in recs if r.get("html_tags"))
    if tags:
        P.append(f"În {tags} fișe titlurile întrebărilor afișează taguri HTML (<br/>, <i>).")
    leak = sum(1 for *_, l in randuri_com if l)
    if leak:
        P.append(f"Întrebarea despre examen (VIII.8) este lipită de rubrica „Ce ați vrea să fie altfel” în {leak} comentarii.")
    blank = collections.Counter(r["blank_pages"] for r in recs if r.get("file") and r["n17"])
    if any(k for k in blank):
        P.append("Fișele completate conțin pagini goale (" + ", ".join(f"{v} {'fișă' if v == 1 else 'fișe'} cu {k}" for k, v in sorted(blank.items()) if k) + ").")
    if nume_corupte:
        P.append(f"{nume_corupte} nume de fișiere din arhivă au diacriticele corupte (Ă, Ț, Î).")
    erori = [r["file"] for r in recs if r.get("error")]
    if erori:
        P.append(f"{len(erori)} fișe PDF nu au putut fi citite: " + ", ".join(erori))
    return rec, {k: v for k, v in prob.items() if v}


# ---------------------------------------------------------------------------
# 6. Exporturi
# ---------------------------------------------------------------------------

def scrie_csv(path: Path, antet, randuri):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(antet)
        w.writerows(randuri)


def exporta(out: Path, data, recs, randuri_com):
    sh = {s["id"]: s for s in data["sheets"]}
    num = lambda x: "" if x is None else str(x).replace(".", ",")
    scrie_csv(out / "fise.csv",
              ["id_fisa", "departament", "program", "tip", "an", "cadru_didactic", "disciplina", "n_semestru", "scor_D1_D7",
               "procent_S_la_D6", "n_sesiune", "scor_D8", "nr_comentarii", "fisa_lipsa", "fisier"],
              [[s["id"], s["dep"], s["p"], s["t"], s["an"], s["cd"], s["d"], s["n"], num(s["s"]), num(s["sS"]), s["n8"], num(s["s8"]),
                s["nc"], "da" if s["miss"] else "", recs[s["id"]].get("file") or ""] for s in data["sheets"]])
    scrie_csv(out / "comentarii.csv",
              ["id_fisa", "program", "tip", "cadru_didactic", "disciplina", "n_semestru", "scor_D1_D7", "rubrica", "comentariu",
               "fara_continut", "teme_auto", "tema_manuala"],
              [[i, sh[i]["p"], sh[i]["t"], sh[i]["cd"], sh[i]["d"], sh[i]["n"], num(sh[i]["s"]), rub, t, "da" if triv else "",
                " | ".join(et), ""] for i, rub, t, triv, et, _ in randuri_com])
    scrie_csv(out / "programe.csv",
              ["cod", "program", "prescurtare", "nivel", "departament", "studenti", "fise", "fise_cu_raspunsuri", "fise_cu_raspunsuri_semestru",
               "raspunsuri_semestru", "raspunsuri_sesiune", "raspunsuri_per_student", "DEv_oficial_C", "DEv_oficial_A",
               "scor", "scor_curs", "scor_aplicatii", "procent_S_la_D6"],
              [[p["code"], p["name"], p["short"], p["niv"], p["depName"], p["stud"], p["sheets"], p["withAny"], p["withSem"], p["sem"], p["exam"],
                num(round(p["sem"] / p["stud"], 2) if p["stud"] else None), num(p["repDEv"][0] if p["repDEv"] else None),
                num(p["repDEv"][1] if p["repDEv"] else None), num(p["score"]), num(p["scoreC"]), num(p["scoreA"]),
                pct(p["d6"]["S"], p["d6"]["n"])] for p in data["progs"]])
    scrie_csv(out / "dimensiuni.csv",
              ["dimensiune", "nume", "tip", "n", "S", "M", "R", "procent_S", "procent_R", "medie_estimata"],
              [[d["id"], d["name"], t, d[t]["n"], d[t]["S"], d[t]["M"], d[t]["R"], pct(d[t]["S"], d[t]["n"]), pct(d[t]["R"], d[t]["n"]),
                num(d[t]["est"])] for d in data["dims"] for t in "CA"])


# ---------------------------------------------------------------------------
# 7. Raportul HTML
# ---------------------------------------------------------------------------

def bloc(text: str, nume: str) -> str:
    m = re.search(rf"<!--\s*{nume}\s*-->(.*?)<!--\s*/{nume}\s*-->", text, re.S)
    return m.group(1).strip() if m else ""


def li_constatare(eticheta, clasa, titlu, text):
    return f'<li><span class="tag {clasa}">{eticheta}</span><p><b>{titlu}</b> {text}</p></li>'


def genereaza_html(data, rec, prob, teme_neg, teme_poz, info_com, observatii: str, an_sem: str, facultate: str, manual: bool) -> str:
    T = data["tot"]
    e = html.escape
    # constatări generate din cifre
    mediana = str(int(T["median"])) if float(T["median"]).is_integer() else f2(T["median"], 1)
    cst = [li_constatare("Participare", "crit",
                         f"Doar {pct(T['withAny'], T['sheets'])}% din cele {T['sheets']} de poziții evaluabile (cadru didactic × disciplină × curs/aplicații) au primit măcar un răspuns.",
                         f"Fișa tipică are {mediana} respondenți; {T['n12']} fișe se bazează pe 1–2 studenți. "
                         + (f"Rata reală de completare este de {f2(T['rateC'], 1)}% la curs și {f2(T['rateA'], 1)}% la aplicații din chestionarele posibile." if T["rateC"] else ""))]
    if T["facDEv"]:
        cst.append(li_constatare("Metodă", "warn", "Raportul facultății supraestimează acoperirea.",
                                 f"DEv% ({round(T['facDEv'][0])}%) și CDEv% ({round(T['facCDEv'][0])}–{round(T['facCDEv'][1])}%) sunt calculate doar față de disciplinele și cadrele care au primit cel puțin un răspuns. "
                                 f"Raportat la toate fișele generate, au primit răspunsuri pe parcursul semestrului {T['covC'][0]}% din discipline la curs și {T['covA'][0]}% la aplicații, "
                                 f"iar {T['teachers'] - T['teachAny']} din {T['teachers']} de cadre didactice nu au primit niciun răspuns."))
    dimA = max(data["dims"][:6], key=lambda d: d["A"]["S"] / max(1, d["A"]["n"]))
    cst.append(li_constatare("Rezultat", "ok", f"Media ponderată D1–D7 este {f2(T['score'])} din 5, iar {T['shareR']}% din aprecieri sunt la nivel ridicat.",
                             f"Cursurile ({f2(T['scoreC'])}) sunt apreciate {'mai bine' if (T['scoreC'] or 0) >= (T['scoreA'] or 0) else 'mai slab'} decât aplicațiile ({f2(T['scoreA'])}). "
                             f"Cea mai slabă dimensiune la aplicații este „{dimA['name']}” ({pct(dimA['A']['S'], dimA['A']['n'])}% nivel scăzut, față de {pct(dimA['C']['S'], dimA['C']['n'])}% la curs)."))
    mari = [p for p in data["progs"] if p["sem"] >= 30 and p["score"] is not None]
    if mari:
        pmin = min(mari, key=lambda p: p["score"])
        sursa = "aplicații" if (pmin["scoreA"] or 5) < (pmin["scoreC"] or 5) else "curs"
        cst.append(li_constatare("Program", "warn", f"{e(pmin['name'])} are cel mai scăzut scor dintre programele cu cel puțin 30 de răspunsuri ({f2(pmin['score'])}), cu {pct(pmin['d6']['S'], pmin['d6']['n'])}% aprecieri globale la nivel scăzut.",
                                 f"Diferența vine din {sursa}: curs {f2(pmin['scoreC'])}, aplicații {f2(pmin['scoreA'])}."))
    cst.append(bloc(observatii, "REZUMAT"))
    nprob = sum(len(v) for v in prob.values())
    cst.append(li_constatare("Date", "info", f"Verificările automate au găsit {nprob} probleme de completare, metodă sau generare a documentelor.", "Detalii în secțiunea „Calitatea datelor”."))
    probleme_html = "".join(f'<div class="card"><h3>{e(k)}</h3><ul class="plain">' + "".join(f"<li>{e(x)}</li>" for x in v) + "</ul></div>" for k, v in prob.items())
    atrib = bloc(observatii, "ATRIBUIRE")
    if atrib:
        probleme_html += f'<div class="card"><h3>Atribuirea evaluărilor</h3><ul class="plain">{atrib}</ul></div>'
    semnalari = bloc(observatii, "SEMNALARI") or '<p class="note">Adăugați semnalările din comentarii în fișierul de observații (blocul SEMNALARI).</p>'
    masuri = bloc(observatii, "MASURI") or '<li><b>Completați măsurile în fișierul de observații (blocul MASURI).</b></li>'
    practici = bloc(observatii, "PRACTICI")
    sablon = (AICI / "sablon_raport.html").read_text(encoding="utf-8")
    vals = {
        "AN_SEM": e(an_sem), "FACULTATE": e(facultate), "CONSTATARI": "\n".join(x for x in cst if x),
        "N_FISE": str(T["sheets"]), "N_GOALE": str(T["empty"]), "N_PDF": str(sum(1 for s in data["sheets"] if not s["miss"])),
        "N_LIPSA": str(sum(1 for s in data["sheets"] if s["miss"])), "N_CU_SEM": str(T["withSem"]),
        "N_EXAM": str(T["exam"]), "N_EXAM1": str(T["exam1"]), "N_CU_EXAM": str(T["withExam"]), "D8": f2(T["d8"]),
        "PROG_FARA_SESIUNE": e(", ".join(T["lowExamProgs"])) or "—",
        "N_SCOR": str(T["withScore"]), "N_PESTE45": str(T["over45"]), "N_LOW_SMALL": str(T["lowSmall"]),
        "PRAG_N": str(data["prag"]["n"]), "PRAG_SCOR": f2(data["prag"]["scor"], 1), "PRAG_S": str(data["prag"]["s"]),
        "N_FISE_COM": str(info_com["fise"]), "N_PLACUT": str(info_com["placut"]), "N_ALTFEL": str(info_com["altfel"]),
        "N_TOT_COM": str(info_com["placut"] + info_com["altfel"]), "N_SUBST_POZ": str(info_com["subst_poz"]), "N_SUBST_NEG": str(info_com["subst_neg"]),
        "NOTA_TEME": "Codificare manuală (coloana tema_manuala din comentarii.csv)." if manual else "Etichetare automată după cuvinte cheie, orientativă. Pentru raportul final, completați coloana tema_manuala din comentarii.csv și rulați din nou cu --teme-manuale.",
        "RECONCILIERI": "".join(f"<li>{e(x)}</li>" for x in rec), "PROBLEME": probleme_html,
        "SEMNALARI": semnalari, "MASURI": masuri, "PRACTICI": practici,
        "TEME_NEG": json.dumps(sorted(teme_neg.items(), key=lambda x: -x[1]), ensure_ascii=False),
        "TEME_POZ": json.dumps(sorted(teme_poz.items(), key=lambda x: -x[1]), ensure_ascii=False),
        "DATA": json.dumps(data, ensure_ascii=False),
    }
    out = sablon
    for k, v in vals.items():
        out = out.replace("{{" + k + "}}", v)
    return out


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description="Analiza evaluării activității didactice de către studenți (UNITBV).")
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--zip", type=Path, help="arhiva exportată din platformă (ex. IESC.zip)")
    src.add_argument("--dir", type=Path, help="folderul cu fișierele deja dezarhivate")
    ap.add_argument("--raport-facultate", type=Path, help="raportul de sinteză al facultății (.docx), opțional")
    ap.add_argument("--out", type=Path, default=Path("rezultate"), help="folderul de ieșire (implicit: rezultate)")
    ap.add_argument("--observatii", type=Path, default=AICI / "observatii.html",
                    help="fișierul cu observațiile calitative scrise de mână (implicit: observatii.html)")
    ap.add_argument("--teme-manuale", type=Path, help="comentarii.csv cu coloana tema_manuala completată")
    ap.add_argument("--an", default="", help="textul pentru anul și semestrul afișate în raport (ex. „2025–2026, semestrul 2”)")
    ap.add_argument("--prag-n", type=int, default=3, help="nr. minim de respondenți pentru lista de cazuri (implicit 3)")
    ap.add_argument("--prag-scor", type=float, default=3.5, help="scor sub care o fișă intră în listă (implicit 3,5)")
    ap.add_argument("--prag-s", type=int, default=50, help="%% aprecieri globale la nivel S peste care o fișă intră în listă (implicit 50)")
    ap.add_argument("--procese", type=int, default=4, help="câte fișe PDF se citesc în paralel (implicit 4)")
    a = ap.parse_args()

    a.out.mkdir(parents=True, exist_ok=True)
    nume_corupte = 0
    if a.zip:
        print(f"Extrag {a.zip} ...")
        root, nume_corupte = extrage_arhiva(a.zip, a.out / "_arhiva")
    else:
        root = a.dir
        nume_corupte = sum(1 for p in root.rglob("*") if repara_mojibake(p.name) != p.name)
    pdfs, rap_paths, sin_paths = gaseste_fisiere(root)
    print(f"Găsite: {len(pdfs)} fișe PDF, {len(rap_paths)} rapoarte pe program, {len(sin_paths)} tabele de sinteză.")
    if not pdfs:
        sys.exit("Nu am găsit nicio fișă PDF. Verificați calea dată la --zip sau --dir.")

    print("Citesc fișele PDF (durează câteva minute) ...")
    with Pool(max(1, a.procese)) as pool:
        fise = []
        for i, r in enumerate(pool.imap(citeste_fisa, [str(p) for p in pdfs], chunksize=4), 1):
            fise.append(r)
            if i % 50 == 0 or i == len(pdfs):
                print(f"  {i}/{len(pdfs)}")
    print("Citesc rapoartele și tabelele de sinteză ...")
    rapoarte = [citeste_raport(p) for p in rap_paths]
    sinteze = [citeste_sinteza(p) for p in sin_paths]
    fac = citeste_raport(a.raport_facultate) if a.raport_facultate else None

    recs, lipsa, fara_sinteza = imbina(fise, rapoarte, sinteze)
    data = analizeaza(recs, rapoarte, sinteze, fac, a.prag_n, a.prag_scor, a.prag_s)
    tneg, tpoz, randuri_com, subst_neg, subst_poz = teme(recs)
    manual = False
    if a.teme_manuale:
        mneg, mpoz = teme_manuale(a.teme_manuale, recs)
        if mneg or mpoz:
            tneg, tpoz, manual = mneg, mpoz, True
    info_com = dict(fise=sum(1 for r in recs if r["likes"] or r["dislikes"]), placut=sum(len(r["likes"]) for r in recs),
                    altfel=sum(len(r["dislikes"]) for r in recs), subst_neg=subst_neg, subst_poz=subst_poz)
    rec, prob = verifica(recs, rapoarte, sinteze, fac, lipsa, fara_sinteza, nume_corupte, randuri_com)

    an_sem = a.an
    if not an_sem:
        hdr = " ".join(t for s in sinteze for t in s["header"])
        m1 = re.search(r"AN UNIVERSITAR\s*([\d\-–/]+)", hdr); m2 = re.search(r"SEMESTRUL:\s*(\d)", hdr)
        an_sem = (f"{m1.group(1).replace('-', '–')}, semestrul {m2.group(1)}" if m1 and m2 else "")
    observatii = a.observatii.read_text(encoding="utf-8") if a.observatii and a.observatii.exists() else ""
    facultate = next((t for s in sinteze for t in s["header"] if t.upper().startswith("FACULTATEA")), "")
    mici = {"DE", "ȘI", "SI", "A", "AL", "IN", "ÎN"}
    facultate = " ".join(w.lower() if w in mici else w.capitalize() for w in facultate.split()).replace(" si ", " și ")

    exporta(a.out, data, recs, randuri_com)
    (a.out / "verificari.txt").write_text("RECONCILIERI\n" + "\n".join("- " + x for x in rec) + "\n\n" +
                                          "\n\n".join(k.upper() + "\n" + "\n".join("- " + x for x in v) for k, v in prob.items()) + "\n", encoding="utf-8")
    json.dump(dict(fise=recs, rapoarte=[{k: v for k, v in r.items() if k != "text"} for r in rapoarte], sinteze=sinteze,
                   facultate={k: v for k, v in (fac or {}).items() if k != "text"}, agregate=data),
              open(a.out / "date_complete.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
    (a.out / "raport_evaluare.html").write_text(
        genereaza_html(data, rec, prob, tneg, tpoz, info_com, observatii, an_sem, facultate, manual), encoding="utf-8")

    T = data["tot"]
    print("\nGata.")
    print(f"  Poziții evaluabile: {T['sheets']} (cu răspunsuri: {T['withAny']}, goale: {T['empty']})")
    print(f"  Răspunsuri: {T['sem']} în semestru, {T['exam']} în sesiune; scor mediu ponderat D1–D7: {f2(T['score'])}")
    print(f"  Probleme de date: {sum(len(v) for v in prob.values())} (vezi verificari.txt)")
    print(f"  Raport: {(a.out / 'raport_evaluare.html').resolve()}")


if __name__ == "__main__":
    main()
